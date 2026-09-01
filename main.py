import tkinter as tk
from tkinter import filedialog, messagebox, ttk
from pathlib import Path
from copy import copy

import pandas as pd
from openpyxl import load_workbook


class ExcelAutomationApp:

    def __init__(self, root):

        self.root = root

        self.root.title("Excel & CSV Automation Tool")
        self.root.geometry("1100x700")
        self.root.minsize(900, 600)

        # =====================================================
        # VARIABLES
        # =====================================================

        self.df = None
        self.file_path = None

        self.header_row = None
        self.data_start_row = None

        self.original_headers = []
        self.original_column_numbers = {}

        self.original_max_row = None
        self.original_max_column = None

        self.original_sheet_name = None

        # =====================================================
        # INTERFAZ
        # =====================================================

        self.setup_style()
        self.create_interface()

    # =========================================================
    # ESTILO
    # =========================================================

    def setup_style(self):

        style = ttk.Style()

        try:
            style.theme_use("clam")
        except Exception:
            pass

        style.configure(
            "Title.TLabel",
            font=("Segoe UI", 26, "bold"),
            foreground="#333333"
        )

        style.configure(
            "Subtitle.TLabel",
            font=("Segoe UI", 12),
            foreground="#666666"
        )

        style.configure(
            "Action.TButton",
            font=("Segoe UI", 11),
            padding=(18, 10)
        )

        style.configure(
            "Treeview",
            font=("Segoe UI", 10),
            rowheight=28
        )

        style.configure(
            "Treeview.Heading",
            font=("Segoe UI", 10, "bold")
        )

    # =========================================================
    # INTERFAZ
    # =========================================================

    def create_interface(self):

        # -----------------------------------------------------
        # HEADER
        # -----------------------------------------------------

        header = ttk.Frame(self.root)

        header.pack(
            fill="x",
            padx=35,
            pady=(25, 10)
        )

        title = ttk.Label(
            header,
            text="Excel & CSV Automation Tool",
            style="Title.TLabel"
        )

        title.pack()

        subtitle = ttk.Label(
            header,
            text="Load, clean, organize and export your data",
            style="Subtitle.TLabel"
        )

        subtitle.pack(
            pady=(5, 0)
        )

        # -----------------------------------------------------
        # BOTONES
        # -----------------------------------------------------

        buttons = ttk.Frame(self.root)

        buttons.pack(
            fill="x",
            padx=35,
            pady=20
        )

        self.load_button = ttk.Button(
            buttons,
            text="Load Excel / CSV",
            style="Action.TButton",
            command=self.load_file
        )

        self.load_button.pack(
            side="left",
            padx=5
        )

        self.clean_button = ttk.Button(
            buttons,
            text="Clean Data",
            style="Action.TButton",
            command=self.clean_data
        )

        self.clean_button.pack(
            side="left",
            padx=5
        )

        self.duplicates_button = ttk.Button(
            buttons,
            text="Remove Duplicates",
            style="Action.TButton",
            command=self.remove_duplicates
        )

        self.duplicates_button.pack(
            side="left",
            padx=5
        )

        self.export_button = ttk.Button(
            buttons,
            text="Export",
            style="Action.TButton",
            command=self.export_file
        )

        self.export_button.pack(
            side="left",
            padx=5
        )

        # -----------------------------------------------------
        # INFORMACIÓN
        # -----------------------------------------------------

        info_frame = ttk.Frame(self.root)

        info_frame.pack(
            fill="x",
            padx=35,
            pady=(0, 10)
        )

        self.file_label = ttk.Label(
            info_frame,
            text="No file loaded",
            font=("Segoe UI", 10)
        )

        self.file_label.pack(
            side="left"
        )

        self.size_label = ttk.Label(
            info_frame,
            text="Rows: 0 | Columns: 0",
            font=("Segoe UI", 10)
        )

        self.size_label.pack(
            side="right"
        )

        # -----------------------------------------------------
        # PREVIEW
        # -----------------------------------------------------

        preview_title = ttk.Label(
            self.root,
            text="Data Preview",
            font=("Segoe UI", 18, "bold")
        )

        preview_title.pack(
            anchor="w",
            padx=35,
            pady=(10, 10)
        )

        table_frame = ttk.Frame(self.root)

        table_frame.pack(
            fill="both",
            expand=True,
            padx=35,
            pady=(0, 30)
        )

        vertical_scroll = ttk.Scrollbar(
            table_frame,
            orient="vertical"
        )

        horizontal_scroll = ttk.Scrollbar(
            table_frame,
            orient="horizontal"
        )

        self.tree = ttk.Treeview(
            table_frame,
            show="headings",
            yscrollcommand=vertical_scroll.set,
            xscrollcommand=horizontal_scroll.set
        )

        vertical_scroll.config(
            command=self.tree.yview
        )

        horizontal_scroll.config(
            command=self.tree.xview
        )

        self.tree.grid(
            row=0,
            column=0,
            sticky="nsew"
        )

        vertical_scroll.grid(
            row=0,
            column=1,
            sticky="ns"
        )

        horizontal_scroll.grid(
            row=1,
            column=0,
            sticky="ew"
        )

        table_frame.rowconfigure(
            0,
            weight=1
        )

        table_frame.columnconfigure(
            0,
            weight=1
        )

    # =========================================================
    # BUSCAR ENCABEZADOS
    # =========================================================

    def find_header_row(self, worksheet):

        for row in worksheet.iter_rows():

            values = []

            for cell in row:

                if cell.value is None:
                    continue

                value = str(
                    cell.value
                ).strip().upper()

                values.append(value)

            has_company = any(
                "COMPA" in value
                for value in values
            )

            has_digitado = any(
                "DIGITADO" in value
                for value in values
            )

            has_bajado = any(
                "BAJADO" in value
                for value in values
            )

            if (
                has_company
                and has_digitado
                and has_bajado
            ):
                return row[0].row

        return None

    # =========================================================
    # LEER ENCABEZADOS ORIGINALES
    # =========================================================

    def read_original_headers(
        self,
        worksheet
    ):

        self.original_headers = []
        self.original_column_numbers = {}

        for cell in worksheet[self.header_row]:

            column_number = cell.column

            value = cell.value

            if value is None:

                header = ""

            else:

                header = str(
                    value
                ).strip()

            self.original_headers.append(
                header
            )

            if header != "":

                self.original_column_numbers[
                    header
                ] = column_number

    # =========================================================
    # CARGAR ARCHIVO
    # =========================================================

    def load_file(self):

        file_path = filedialog.askopenfilename(
            title="Select Excel or CSV file",
            filetypes=[
                ("Excel files", "*.xlsx"),
                ("CSV files", "*.csv"),
                ("All files", "*.*")
            ]
        )

        if not file_path:
            return

        try:

            extension = Path(
                file_path
            ).suffix.lower()

            # =================================================
            # CSV
            # =================================================

            if extension == ".csv":

                self.df = pd.read_csv(
                    file_path,
                    dtype=str
                )

                self.df = self.df.fillna("")

                self.file_path = file_path

                self.header_row = 1
                self.data_start_row = 2

                self.original_sheet_name = None

            # =================================================
            # EXCEL
            # =================================================

            elif extension == ".xlsx":

                workbook = load_workbook(
                    file_path,
                    data_only=False
                )

                worksheet = workbook.active

                self.original_sheet_name = (
                    worksheet.title
                )

                self.header_row = (
                    self.find_header_row(
                        worksheet
                    )
                )

                if self.header_row is None:

                    messagebox.showerror(
                        "Error",
                        "Could not find the Excel headers."
                    )

                    return

                self.data_start_row = (
                    self.header_row + 1
                )

                self.original_max_row = (
                    worksheet.max_row
                )

                self.original_max_column = (
                    worksheet.max_column
                )

                self.read_original_headers(
                    worksheet
                )

                # ---------------------------------------------
                # LEER DATOS
                # ---------------------------------------------

                data = []

                for row_number in range(
                    self.data_start_row,
                    worksheet.max_row + 1
                ):

                    values = []

                    has_data = False

                    for column_number in range(
                        1,
                        worksheet.max_column + 1
                    ):

                        cell = worksheet.cell(
                            row=row_number,
                            column=column_number
                        )

                        value = cell.value

                        if value is not None:

                            has_data = True

                        values.append(
                            "" if value is None
                            else value
                        )

                    if has_data:

                        data.append(values)

                # ---------------------------------------------
                # CREAR HEADERS
                # ---------------------------------------------

                headers = []

                for index, header in enumerate(
                    self.original_headers
                ):

                    if header == "":

                        headers.append(
                            f"Column {index + 1}"
                        )

                    else:

                        headers.append(
                            header
                        )

                self.df = pd.DataFrame(
                    data,
                    columns=headers
                )

                self.df = self.df.fillna("")

                self.file_path = file_path

            else:

                messagebox.showerror(
                    "Error",
                    "Unsupported file type."
                )

                return

            self.update_preview()

            messagebox.showinfo(
                "File Loaded",
                "File loaded successfully.\n\n"
                f"Rows: {len(self.df)}\n"
                f"Columns: {len(self.df.columns)}"
            )

        except Exception as e:

            messagebox.showerror(
                "Error loading file",
                str(e)
            )

    # =========================================================
    # PREVIEW
    # =========================================================

    def update_preview(self):

        for item in self.tree.get_children():

            self.tree.delete(item)

        self.tree["columns"] = []

        if self.df is None:

            return

        columns = list(
            self.df.columns
        )

        self.tree["columns"] = columns

        for column in columns:

            name = str(column)

            self.tree.heading(
                name,
                text=name
            )

            self.tree.column(
                name,
                width=180,
                minwidth=100,
                anchor="w"
            )

        preview = self.df.head(500)

        for _, row in preview.iterrows():

            values = []

            for column in columns:

                value = row[column]

                if pd.isna(value):

                    value = ""

                values.append(
                    str(value)
                )

            self.tree.insert(
                "",
                "end",
                values=values
            )

        if self.file_path:

            self.file_label.config(
                text=(
                    f"Loaded: "
                    f"{Path(self.file_path).name}"
                )
            )

        self.size_label.config(
            text=(
                f"Rows: {len(self.df)} | "
                f"Columns: {len(self.df.columns)}"
            )
        )

        self.auto_resize_columns()

    # =========================================================
    # AJUSTAR COLUMNAS PREVIEW
    # =========================================================

    def auto_resize_columns(self):

        if self.df is None:
            return

        for column in self.df.columns:

            max_length = len(
                str(column)
            )

            for value in self.df[column].head(100):

                length = len(
                    str(value)
                )

                if length > max_length:

                    max_length = length

            width = max(
                100,
                min(
                    max_length * 8,
                    350
                )
            )

            self.tree.column(
                str(column),
                width=width
            )

    # =========================================================
    # CLEAN DATA
    # =========================================================

    def clean_data(self):

        if self.df is None:

            messagebox.showwarning(
                "No data",
                "Please load an Excel or CSV file first."
            )

            return

        try:

            before = len(
                self.df
            )

            # ---------------------------------------------
            # LIMPIAR ESPACIOS
            # ---------------------------------------------

            for column in self.df.columns:

                self.df[column] = (
                    self.df[column]
                    .fillna("")
                    .astype(str)
                    .str.strip()
                )

            # ---------------------------------------------
            # ELIMINAR FILAS COMPLETAMENTE VACÍAS
            # ---------------------------------------------

            self.df = self.df[
                self.df.astype(str)
                .apply(
                    lambda row:
                    row.str.strip().ne("").any(),
                    axis=1
                )
            ]

            self.df.reset_index(
                drop=True,
                inplace=True
            )

            removed = (
                before -
                len(self.df)
            )

            self.update_preview()

            messagebox.showinfo(
                "Clean Data",
                "Data cleaned successfully.\n\n"
                f"Empty rows removed: {removed}\n"
                f"Rows: {len(self.df)}\n"
                f"Columns: {len(self.df.columns)}"
            )

        except Exception as e:

            messagebox.showerror(
                "Cleaning error",
                str(e)
            )

    # =========================================================
    # DUPLICADOS
    # =========================================================

    def remove_duplicates(self):

        if self.df is None:

            messagebox.showwarning(
                "No data",
                "Please load an Excel or CSV file first."
            )

            return

        try:

            before = len(
                self.df
            )

            self.df = self.df.drop_duplicates(
                keep="first"
            )

            self.df.reset_index(
                drop=True,
                inplace=True
            )

            removed = (
                before -
                len(self.df)
            )

            self.update_preview()

            messagebox.showinfo(
                "Remove Duplicates",
                "Duplicate removal completed.\n\n"
                f"Duplicates removed: {removed}\n"
                f"Remaining rows: {len(self.df)}"
            )

        except Exception as e:

            messagebox.showerror(
                "Duplicate error",
                str(e)
            )

    # =========================================================
    # COPIAR ESTILO DE UNA CELDA
    # =========================================================

    def copy_cell_style(
        self,
        source_cell,
        target_cell
    ):

        target_cell.font = copy(
            source_cell.font
        )

        target_cell.fill = copy(
            source_cell.fill
        )

        target_cell.border = copy(
            source_cell.border
        )

        target_cell.alignment = copy(
            source_cell.alignment
        )

        target_cell.number_format = (
            source_cell.number_format
        )

        target_cell.protection = copy(
            source_cell.protection
        )

    # =========================================================
    # EXPORTAR EXCEL
    # =========================================================

    def export_excel(
        self,
        save_path
    ):

        # =====================================================
        # ABRIR ORIGINAL
        # =====================================================

        workbook = load_workbook(
            self.file_path,
            data_only=False
        )

        worksheet = workbook[
            self.original_sheet_name
        ]

        # =====================================================
        # VOLVER A ENCONTRAR HEADER
        # =====================================================

        header_row = (
            self.find_header_row(
                worksheet
            )
        )

        if header_row is None:

            raise Exception(
                "Could not find original header row."
            )

        data_start_row = (
            header_row + 1
        )

        # =====================================================
        # LEER COLUMNAS
        # =====================================================

        column_map = {}

        for cell in worksheet[header_row]:

            if cell.value is None:
                continue

            name = str(
                cell.value
            ).strip()

            if name:

                column_map[name] = (
                    cell.column
                )

        # =====================================================
        # IMPORTANTE:
        #
        # NO TOCAMOS NADA ANTES DEL HEADER.
        #
        # Esto incluye:
        #
        # - Título
        # - Logos
        # - Celdas combinadas
        # - Colores
        # - Bordes
        # - Filas superiores
        # =====================================================

        # =====================================================
        # GUARDAR ESTILOS DE TODAS LAS FILAS DE DATOS
        # =====================================================

        styles = {}

        for row_number in range(
            data_start_row,
            worksheet.max_row + 1
        ):

            styles[row_number] = {}

            for column_number in range(
                1,
                worksheet.max_column + 1
            ):

                cell = worksheet.cell(
                    row=row_number,
                    column=column_number
                )

                styles[row_number][
                    column_number
                ] = {
                    "font": copy(cell.font),
                    "fill": copy(cell.fill),
                    "border": copy(cell.border),
                    "alignment": copy(cell.alignment),
                    "number_format": cell.number_format,
                    "protection": copy(cell.protection)
                }

        # =====================================================
        # LIMPIAR SOLAMENTE CONTENIDO DE DATOS
        #
        # NUNCA TOCAR HEADER NI TÍTULO
        # =====================================================

        for row_number in range(
            data_start_row,
            worksheet.max_row + 1
        ):

            for column_number in range(
                1,
                worksheet.max_column + 1
            ):

                worksheet.cell(
                    row=row_number,
                    column=column_number
                ).value = None

        # =====================================================
        # ESCRIBIR DATAFRAME
        # =====================================================

        for df_index, (_, row) in enumerate(
            self.df.iterrows()
        ):

            target_row = (
                data_start_row +
                df_index
            )

            # -------------------------------------------------
            # Si tenemos estilo de una fila original,
            # reutilizarlo.
            # -------------------------------------------------

            if target_row in styles:

                source_styles = styles[
                    target_row
                ]

            else:

                # Usar la última fila con estilo
                if styles:

                    last_style_row = max(
                        styles.keys()
                    )

                    source_styles = styles[
                        last_style_row
                    ]

                else:

                    source_styles = {}

            # -------------------------------------------------
            # ESCRIBIR CADA COLUMNA
            # -------------------------------------------------

            for column_index, column_name in enumerate(
                self.df.columns
            ):

                value = row[column_name]

                # ---------------------------------------------
                # Encontrar columna original
                # ---------------------------------------------

                if column_name in column_map:

                    target_column = (
                        column_map[
                            column_name
                        ]
                    )

                else:

                    # Buscar por posición
                    target_column = (
                        column_index + 1
                    )

                cell = worksheet.cell(
                    row=target_row,
                    column=target_column
                )

                # ---------------------------------------------
                # VALOR
                # ---------------------------------------------

                if pd.isna(value):

                    cell.value = None

                else:

                    cell.value = value

                # ---------------------------------------------
                # FORMATO
                # ---------------------------------------------

                if target_column in source_styles:

                    style_data = (
                        source_styles[
                            target_column
                        ]
                    )

                    cell.font = copy(
                        style_data["font"]
                    )

                    cell.fill = copy(
                        style_data["fill"]
                    )

                    cell.border = copy(
                        style_data["border"]
                    )

                    cell.alignment = copy(
                        style_data["alignment"]
                    )

                    cell.number_format = (
                        style_data[
                            "number_format"
                        ]
                    )

                    cell.protection = copy(
                        style_data[
                            "protection"
                        ]
                    )

        # =====================================================
        # MUY IMPORTANTE:
        #
        # NO CAMBIAMOS:
        #
        # - merged_cells
        # - row heights
        # - column widths
        # - filtros existentes
        # - colores
        # - título
        # - encabezados
        # - congelación original
        # =====================================================

        # =====================================================
        # GUARDAR
        # =====================================================

        workbook.save(
            save_path
        )

    # =========================================================
    # EXPORTAR CSV
    # =========================================================

    def export_csv(
        self,
        save_path
    ):

        self.df.to_csv(
            save_path,
            index=False,
            encoding="utf-8-sig"
        )

    # =========================================================
    # EXPORT
    # =========================================================

    def export_file(self):

        if self.df is None:

            messagebox.showwarning(
                "No data",
                "Please load an Excel or CSV file first."
            )

            return

        if self.df.empty:

            messagebox.showwarning(
                "No data",
                "There is no data to export."
            )

            return

        try:

            save_path = filedialog.asksaveasfilename(
                title="Export file",
                defaultextension=".xlsx",
                filetypes=[
                    ("Excel file", "*.xlsx"),
                    ("CSV file", "*.csv")
                ]
            )

            if not save_path:
                return

            extension = Path(
                save_path
            ).suffix.lower()

            if extension == ".csv":

                self.export_csv(
                    save_path
                )

            else:

                if not self.file_path:

                    raise Exception(
                        "No original file loaded."
                    )

                self.export_excel(
                    save_path
                )

            messagebox.showinfo(
                "Export",
                "File exported successfully!\n\n"
                f"{save_path}"
            )

        except Exception as e:

            messagebox.showerror(
                "Export error",
                str(e)
            )


# =============================================================
# MAIN
# =============================================================

def main():

    root = tk.Tk()

    app = ExcelAutomationApp(
        root
    )

    root.mainloop()


if __name__ == "__main__":
    main()
    